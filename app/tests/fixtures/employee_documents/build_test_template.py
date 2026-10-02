"""Build synthetic DOCX templates for employee-document tests."""

from io import BytesIO

from docx import Document


def build_template(*, split_placeholder: bool = False) -> bytes:
    document = Document()
    paragraph = document.add_paragraph("Contrato de trabajo de ")
    if split_placeholder:
        paragraph.add_run("{{ employee_")
        paragraph.add_run("full_name }}")
    else:
        paragraph.add_run("{{ employee_full_name }}")

    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run(
        "Salario: {{ contract.monthly_wage }}"
    )

    header = document.sections[0].header.paragraphs[0]
    header.add_run("ASIATI · {{ employee.job_title }}")

    footer = document.sections[0].footer.paragraphs[0]
    footer.add_run("Documento {{ employee_document_number }}")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def extract_all_text(docx_bytes: bytes) -> str:
    document = Document(BytesIO(docx_bytes))
    parts = []
    parts.extend(paragraph.text for paragraph in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    for section in document.sections:
        parts.extend(paragraph.text for paragraph in section.header.paragraphs)
        parts.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return "\n".join(parts)
