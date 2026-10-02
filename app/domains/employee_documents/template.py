"""Controlled DOCX template manifest and renderer."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import json
import re

from docx import Document


_ALLOWED_SOURCES = {
    "employee.full_name",
    "employee.job_title",
    "employee.department",
    "employee.email",
    "contract.contract_type",
    "contract.start_date",
    "contract.end_date",
    "contract.monthly_wage",
    "manual",
}
_PLACEHOLDER_RE = re.compile(r"{{\s*([^{}]+?)\s*}}")


class TemplateValidationError(ValueError):
    pass


class MissingTemplateVariables(TemplateValidationError):
    def __init__(self, fields: list[str]):
        self.fields = sorted(set(fields))
        super().__init__(
            "Faltan variables obligatorias de plantilla: "
            + ", ".join(self.fields)
        )


@dataclass(frozen=True)
class TemplateField:
    name: str
    label: str
    source: str
    required: bool


@dataclass(frozen=True)
class TemplateManifest:
    template_version: str
    document_type: str
    fields: tuple[TemplateField, ...]

    @classmethod
    def from_json_bytes(cls, data: bytes) -> "TemplateManifest":
        try:
            raw = json.loads(data.decode("utf-8"))
        except Exception as exc:
            raise TemplateValidationError(
                "El manifest de la plantilla no es JSON válido."
            ) from exc

        if not isinstance(raw, dict):
            raise TemplateValidationError("El manifest debe ser un objeto JSON.")
        version = str(raw.get("template_version") or "").strip()
        document_type = str(raw.get("document_type") or "").strip().upper()
        fields_raw = raw.get("fields")
        if not version or document_type != "CONTRACT" or not isinstance(fields_raw, list):
            raise TemplateValidationError("El manifest de contrato está incompleto.")

        fields: list[TemplateField] = []
        names: set[str] = set()
        for item in fields_raw:
            if not isinstance(item, dict):
                raise TemplateValidationError("Campo de manifest inválido.")
            name = str(item.get("name") or "").strip()
            label = str(item.get("label") or "").strip()
            source = str(item.get("source") or "").strip()
            required = bool(item.get("required", False))
            if not name or not label or source not in _ALLOWED_SOURCES:
                raise TemplateValidationError(
                    f"Definición inválida para el campo {name or '<sin nombre>'}."
                )
            if name in names:
                raise TemplateValidationError(
                    f"El campo de plantilla {name} está duplicado."
                )
            names.add(name)
            fields.append(
                TemplateField(
                    name=name,
                    label=label,
                    source=source,
                    required=required,
                )
            )
        if not fields:
            raise TemplateValidationError("El manifest no define campos.")
        return cls(
            template_version=version,
            document_type=document_type,
            fields=tuple(fields),
        )


def _iter_cell_paragraphs(cell):
    for paragraph in cell.paragraphs:
        yield paragraph
    for table in cell.tables:
        for row in table.rows:
            for nested_cell in row.cells:
                yield from _iter_cell_paragraphs(nested_cell)


def _iter_all_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_cell_paragraphs(cell)
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            yield paragraph
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from _iter_cell_paragraphs(cell)
        for paragraph in section.footer.paragraphs:
            yield paragraph
        for table in section.footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from _iter_cell_paragraphs(cell)


def _replace_in_paragraph(paragraph, *, values: dict[str, str], declared: set[str]):
    visible = paragraph.text
    matches = list(_PLACEHOLDER_RE.finditer(visible))
    for match in matches:
        name = match.group(1).strip()
        if name not in declared:
            raise TemplateValidationError(
                f"La plantilla contiene una variable no declarada: {name}."
            )
        literal = match.group(0)
        containing = [run for run in paragraph.runs if literal in run.text]
        if not containing:
            raise TemplateValidationError(
                f"La variable {name} está dividida entre runs de Word."
            )
        replacement = str(values.get(name, ""))
        for run in containing:
            run.text = run.text.replace(literal, replacement)


def render_docx(
    template_bytes: bytes,
    *,
    variables: dict[str, str],
    manifest: TemplateManifest,
) -> bytes:
    declared = {field.name for field in manifest.fields}
    missing = [
        field.name
        for field in manifest.fields
        if field.required
        and (
            field.name not in variables
            or variables[field.name] is None
            or not str(variables[field.name]).strip()
        )
    ]
    if missing:
        raise MissingTemplateVariables(missing)

    try:
        document = Document(BytesIO(template_bytes))
    except Exception as exc:
        raise TemplateValidationError(
            "La plantilla DOCX no se puede abrir."
        ) from exc

    for paragraph in _iter_all_paragraphs(document):
        _replace_in_paragraph(
            paragraph,
            values={key: "" if value is None else str(value) for key, value in variables.items()},
            declared=declared,
        )

    unresolved: set[str] = set()
    for paragraph in _iter_all_paragraphs(document):
        unresolved.update(
            match.group(1).strip()
            for match in _PLACEHOLDER_RE.finditer(paragraph.text)
        )
    if unresolved:
        raise TemplateValidationError(
            "La plantilla conserva variables sin resolver: "
            + ", ".join(sorted(unresolved))
        )

    output = BytesIO()
    document.save(output)
    return output.getvalue()
